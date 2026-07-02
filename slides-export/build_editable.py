# Rebuild Module 13 deck as native editable PowerPoint (converts cleanly to Google Slides)
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from content import S

# ---- design tokens ----
TERRA = RGBColor(0xC4, 0x84, 0x6B)
TERRA_DK = RGBColor(0xB5, 0x74, 0x5B)
TERRA_SOFT = RGBColor(0xE8, 0xD5, 0xCA)     # note pill on cream
TERRA_SOFT2 = RGBColor(0xE9, 0xD6, 0xCC)    # prompt pill / header pill soft
TERRA_INNER = RGBColor(0xCE, 0x9E, 0x89)    # inner card on key-idea
TERRA_PILL = RGBColor(0xD2, 0xA4, 0x8F)     # KEY IDEA pill
CREAM = RGBColor(0xEB, 0xE4, 0xDE)
CARD = RGBColor(0xF6, 0xF2, 0xEF)
MINI = RGBColor(0xFC, 0xFA, 0xF9)
INK = RGBColor(0x17, 0x17, 0x1B)
GRAY = RGBColor(0x4E, 0x4E, 0x55)
W_SOFT = RGBColor(0xF5, 0xE7, 0xDF)         # soft white text on terracotta
LINE_LT = RGBColor(0xD7, 0xB2, 0x9F)        # hairline on cream
LINE_W = RGBColor(0xE0, 0xC0, 0xB2)         # hairline on terracotta
QUOTE_MARK = RGBColor(0xD8, 0xA9, 0x94)
CONNECT = RGBColor(0xDD, 0xC0, 0xB2)

FONT = "Montserrat"
PW, PH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(PW)
prs.slide_height = Inches(PH)
BLANK = prs.slide_layouts[6]


def _set_spc(run, val):
    run.font._rPr.set('spc', str(val))


def rect(slide, x, y, w, h, fill, line=None, line_w=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None, shadow=False):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w or 1.0)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        try:
            frac = max(0.0, min(0.5, radius / min(w, h)))
            sp.adjustments[0] = frac
        except Exception:
            pass
    sp.shadow.inherit = False
    st = sp.element.find(qn('p:style'))
    if st is not None:
        sp.element.remove(st)
    return sp


def text(slide, x, y, w, h, runs_lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True, line_spacing=1.0, space_after=0):
    """runs_lines: list of paragraphs; each paragraph = list of (text, size, bold, color, spc, italic)"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        for (t, size, bold, color, spc, *rest) in para:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            if rest and rest[0]:
                r.font.italic = True
            if spc:
                _set_spc(r, spc)
    return tb


def est_lines(t, width_in, size_pt, bold=False):
    cpl = (width_in * 72.0) / (size_pt * (0.60 if bold else 0.55))
    return max(1, math.ceil(len(t) / cpl))


def bg(slide, color):
    rect(slide, -0.03, -0.03, PW + 0.06, PH + 0.06, color, shape=MSO_SHAPE.RECTANGLE)


def footer(slide, dark=True):
    lc = LINE_LT if dark else LINE_W
    tc = INK if dark else RGBColor(0xFF, 0xFF, 0xFF)
    rect(slide, (PW - 3.2) / 2, 6.66, 3.2, 0.012, lc, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 0, 6.82, PW, 0.35, [[("L U K E   H A W K I N S", 10, True, tc, 300)]],
         align=PP_ALIGN.CENTER)


def kicker_rule(slide, k, dark=True):
    tc = INK if dark else W_SOFT
    lc = LINE_LT if dark else LINE_W
    y = 0.40 if dark else 0.60
    text(slide, 0.94, y, 8.0, 0.35, [[(f"MODULE 13   |   VIDEO {k}", 10, True, tc, 300)]])
    rect(slide, 0.94, y + 0.44, 11.45, 0.014, lc, shape=MSO_SHAPE.RECTANGLE)


def heading(slide, h, sub=None):
    size = 30
    if len(h) > 42:
        size = 26
    if len(h) > 52:
        size = 23
    text(slide, 0.94, 1.02, 11.5, 0.85, [[(h, size, True, INK, 0)]])
    if sub:
        text(slide, 0.96, 1.72, 11.4, 0.35, [[(sub, 12, False, GRAY, 0)]])
        return 2.18
    return 2.02


def note_pill(slide, x, y, w, lead, rest, two_line=False, size=10.5):
    h = 0.72 if two_line else 0.5
    rect(slide, x, y, w, h, TERRA_SOFT, radius=0.16)
    runs = []
    if lead:
        runs.append((lead + " ", size, True, INK, 0))
    runs.append((rest, size, False, INK, 0))
    text(slide, x + 0.3, y, w - 0.6, h, [runs], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    return h


def big_card(slide, y0, h):
    rect(slide, 1.14, y0, 11.25, h, CARD, radius=0.22)
    rect(slide, 0.98, y0, 0.09, h, TERRA, radius=0.045)


def marker(slide, x, y, s=0.16):
    rect(slide, x, y, s, s, TERRA, radius=0.05)


def num_circle(slide, x, y, n, d=0.34):
    sp = rect(slide, x, y, d, d, TERRA, shape=MSO_SHAPE.OVAL)
    tf = sp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.name = FONT
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def mini_card(slide, x, y, w, h, title, desc, tsize=13, dsize=10.5, tc=INK, pad=0.24, title_marker=True):
    rect(slide, x, y, w, h, MINI, radius=0.14)
    rect(slide, x + 0.12, y - 0.028, w - 0.24, 0.075, TERRA, radius=0.037)
    tx = x + pad
    tw = w - 2 * pad
    if title_marker:
        marker(slide, x + pad, y + 0.27, 0.14)
        tx = x + pad + 0.26
        tw = w - 2 * pad - 0.26
    paras = [[(title, tsize, True, tc, 0)]]
    if desc:
        paras.append([(desc, dsize, False, GRAY, 0)])
    text(slide, tx, y + 0.16, tw, h - 0.3, paras, line_spacing=1.08, space_after=4)


# ---------- templates ----------

def t_title(sl, d, final=False):
    bg(sl, TERRA)
    rect(sl, 0.38, 0.36, PW - 0.76, PH - 0.72, None, line=LINE_W, line_w=1.0, radius=0.25)
    kicker_rule(sl, d["k"], dark=False)
    lines = d["h"].split("\n")
    paras = [[(ln, 40, True, RGBColor(0xFF, 0xFF, 0xFF), 0)] for ln in lines]
    text(sl, 1.2, 2.3 if not final else 2.35, PW - 2.4, 1.9, paras, align=PP_ALIGN.CENTER, line_spacing=1.05)
    yy = 4.15 if len(lines) > 1 or len(d["h"]) > 38 else 3.6
    if final:
        rect(sl, (PW - 5.0) / 2, yy, 5.0, 0.016, LINE_W, shape=MSO_SHAPE.RECTANGLE)
        yy += 0.35
    text(sl, 2.0, yy, PW - 4.0, 0.5, [[(d["sub"], 16, False, W_SOFT, 0)]], align=PP_ALIGN.CENTER)
    footer(sl, dark=False)


def t_quote(sl, d):
    bg(sl, TERRA)
    rect(sl, 0.38, 0.36, PW - 0.76, PH - 0.72, None, line=LINE_W, line_w=1.0, radius=0.25)
    kicker_rule(sl, d["k"], dark=False)
    text(sl, 1.1, 1.35, 2.0, 1.4, [[("“", 110, True, QUOTE_MARK, 0)]], wrap=False)
    main_size = 28 if len(d["main"]) < 90 else 24
    text(sl, 1.9, 2.45, PW - 3.8, 1.8, [[(d["main"], main_size, True, RGBColor(0xFF, 0xFF, 0xFF), 0)]],
         align=PP_ALIGN.CENTER, line_spacing=1.1)
    if d.get("sub"):
        ml = est_lines(d["main"], PW - 3.8, main_size, True)
        sub_y = 2.45 + ml * (main_size * 1.25 / 72.0) + 0.3
        text(sl, 2.1, sub_y, PW - 4.2, 0.9, [[(d["sub"], 15, False, W_SOFT, 0)]],
             align=PP_ALIGN.CENTER, line_spacing=1.15)
    text(sl, PW - 2.6, 4.9, 1.6, 1.3, [[("”", 110, True, QUOTE_MARK, 0)]], wrap=False, align=PP_ALIGN.RIGHT)
    footer(sl, dark=False)


def t_bullets(sl, d):
    bg(sl, CREAM)
    kicker_rule(sl, d["k"])
    top = heading(sl, d["h"], d.get("sub"))
    footer(sl)
    card_h = 6.42 - top
    big_card(sl, top, card_h)
    n = len(d["items"])
    numbered = d.get("numbered", False)
    note = d.get("note")
    inner_top = top + 0.22
    inner_bot = top + card_h - 0.2
    if note:
        inner_bot -= 0.62
    area = inner_bot - inner_top
    slot = area / n
    if n <= 4:
        ts, ds = 15, 11.5
    elif n == 5:
        ts, ds = 14, 11
    elif n == 6:
        ts, ds = 13, 10.5
    else:
        ts, ds = 12, 10
    x_text = 2.02 if numbered else 1.88
    w_text = 12.1 - x_text
    for i, (title, desc) in enumerate(d["items"]):
        y = inner_top + i * slot
        dl = est_lines(desc, w_text, ds)
        if numbered:
            num_circle(sl, 1.5, y + 0.02, i + 1)
        else:
            marker(sl, 1.56, y + 0.06)
        text(sl, x_text, y, w_text, slot,
             [[(title, ts, True, INK, 0)], [(desc, ds, False, GRAY, 0)]],
             line_spacing=1.05, space_after=2)
    if note:
        note_pill(sl, 1.45, inner_bot + 0.12, 10.65, note[0], note[1])


def t_keyidea(sl, d):
    bg(sl, CREAM)
    kicker_rule(sl, d["k"])
    top = heading(sl, d["h"])
    footer(sl)
    rect(sl, 1.14, top + 0.1, 11.05, 4.05, TERRA, radius=0.24)
    # pill
    rect(sl, 1.55, top + 0.42, 1.5, 0.44, TERRA_PILL, radius=0.22)
    text(sl, 1.55, top + 0.42, 1.5, 0.44, [[("KEY IDEA", 10.5, True, RGBColor(0xFF, 0xFF, 0xFF), 250)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(sl, 1.62, top + 1.08, 10.2, 0.5, [[(d["main"], 20, True, RGBColor(0xFF, 0xFF, 0xFF), 0)]])
    text(sl, 1.62, top + 1.58, 10.2, 0.4, [[(d["desc"], 12.5, False, W_SOFT, 0)]])
    # inner card
    iy = top + 2.12
    rect(sl, 1.55, iy, 10.25, 1.75, TERRA_INNER, radius=0.16)
    for i, b in enumerate(d["bullets"]):
        yy = iy + 0.22 + i * 0.48
        rect(sl, 1.85, yy + 0.08, 0.12, 0.12, RGBColor(0xFF, 0xFF, 0xFF), shape=MSO_SHAPE.OVAL)
        text(sl, 2.12, yy, 9.4, 0.45, [[(b, 12, False, RGBColor(0xFF, 0xFF, 0xFF), 0)]])


def t_grid(sl, d):
    bg(sl, CREAM)
    kicker_rule(sl, d["k"])
    top = heading(sl, d["h"], d.get("sub"))
    footer(sl)
    note = d.get("note")
    prompt = d.get("prompt")
    callout = d.get("callout")
    x0, x1 = 1.15, 12.4
    gap = 0.3
    y = top
    if prompt:
        rect(sl, x0, y, x1 - x0, 0.78, TERRA_SOFT2, radius=0.16)
        text(sl, x0 + 0.35, y + 0.1, x1 - x0 - 0.7, 0.62,
             [[(prompt[0], 10, True, INK, 250)], [(prompt[1], 13, True, INK, 0)]], line_spacing=1.1)
        y += 0.95
    bot = 6.42
    if note:
        bot -= 0.64
    items = d["items"]
    single = d.get("cols3") and len(items) == 3
    rows = 1 if single else 2
    cols = 3
    cw = (x1 - x0 - (cols - 1) * gap) / cols
    if single:
        ch = 1.55
    else:
        ch = (bot - y - 0.35) / 2
    for i, (title, desc) in enumerate(items):
        r, c = divmod(i, cols)
        cx = x0 + c * (cw + gap)
        cy = y + r * (ch + 0.32)
        big_title = est_lines(title, cw - 0.75, 12.5, True) > 1
        mini_card(sl, cx, cy, cw, ch, title, desc, tsize=12.5, dsize=10)
    if single and callout:
        cy = y + ch + 0.35
        wch = bot - cy
        rect(sl, x0 + 1.6, cy, x1 - x0 - 3.2, wch, MINI, radius=0.14)
        rect(sl, x0 + 1.72, cy - 0.028, x1 - x0 - 3.44, 0.075, TERRA, radius=0.037)
        text(sl, x0 + 1.95, cy + 0.16, x1 - x0 - 3.9, wch - 0.3,
             [[(callout[0], 12.5, True, INK, 0)], [(callout[1], 10, False, GRAY, 0)]],
             line_spacing=1.1, space_after=3)
    if note:
        note_pill(sl, x0, 6.42 - 0.52, x1 - x0, note[0], note[1],
                  two_line=est_lines((note[0] or "") + note[1], x1 - x0 - 0.6, 10.5) > 1)


def t_compare(sl, d):
    bg(sl, CREAM)
    kicker_rule(sl, d["k"])
    top = heading(sl, d["h"])
    footer(sl)
    x0 = 1.15
    colw = 5.5
    x_r = x0 + colw + 0.25
    hh = 0.6
    for x, label in ((x0, d["lh"]), (x_r, d["rh"])):
        rect(sl, x, top, colw, hh, TERRA, radius=0.14)
        text(sl, x + 0.3, top, colw - 0.6, hh, [[(label, 13.5, True, RGBColor(0xFF, 0xFF, 0xFF), 0)]],
             anchor=MSO_ANCHOR.MIDDLE)
    body_y = top + hh + 0.15
    body_h = 6.42 - body_y
    for x, items in ((x0, d["left"]), (x_r, d["right"])):
        rect(sl, x, body_y, colw, body_h, MINI, radius=0.16)
        n = len(items)
        slot = (body_h - 0.35) / n
        ts = 12.5 if n <= 4 else 11.5
        ds = 10 if n <= 4 else 9.5
        for i, (title, desc) in enumerate(items):
            y = body_y + 0.2 + i * slot
            marker(sl, x + 0.28, y + 0.05, 0.13)
            text(sl, x + 0.55, y, colw - 0.85, slot,
                 [[(title, ts, True, INK, 0)], [(desc, ds, False, GRAY, 0)]],
                 line_spacing=1.03, space_after=2)


def t_table(sl, d):
    bg(sl, CREAM)
    kicker_rule(sl, d["k"])
    top = heading(sl, d["h"])
    footer(sl)
    x0, x1 = 1.15, 12.4
    c1x, c1w = x0 + 0.35, 4.1
    c2x, c2w = 6.05, 2.1
    c3x, c3w = 8.45, 3.7
    hh = 0.58
    rect(sl, x0, top, x1 - x0, hh, TERRA, radius=0.14)
    for cx, cw_, label, al in ((c1x, c1w, d["cols"][0], PP_ALIGN.LEFT),
                               (c2x, c2w, d["cols"][1], PP_ALIGN.CENTER),
                               (c3x, c3w, d["cols"][2], PP_ALIGN.LEFT)):
        text(sl, cx, top, cw_, hh, [[(label, 11.5, True, RGBColor(0xFF, 0xFF, 0xFF), 200)]],
             align=al, anchor=MSO_ANCHOR.MIDDLE)
    note = d.get("note")
    bot = 6.42 - (0.64 if note else 0)
    rows = d["rows"]
    gap = 0.14
    rh = (bot - top - hh - 0.15 - (len(rows) - 1) * gap) / len(rows)
    y = top + hh + 0.15
    for (t1, sub1, val, desc) in rows:
        rect(sl, x0, y, x1 - x0, rh, MINI, radius=0.12)
        marker(sl, c1x, y + rh / 2 - 0.2, 0.15)
        text(sl, c1x + 0.3, y, c1w - 0.3, rh,
             [[(t1, 13, True, INK, 0)], [(sub1, 9.5, False, GRAY, 0)]],
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1, space_after=2)
        vsize = 20 if len(val) <= 8 else (13 if len(val) <= 12 else 11)
        text(sl, c2x, y, c2w, rh, [[(val, vsize, True, INK, 0)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
        text(sl, c3x, y, c3w, rh, [[(desc, 10, False, GRAY, 0)]],
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
        y += rh + gap
    if note:
        note_pill(sl, x0, 6.42 - 0.52, x1 - x0, note[0], note[1])


def t_timeline(sl, d):
    bg(sl, CREAM)
    kicker_rule(sl, d["k"])
    top = heading(sl, d["h"])
    footer(sl)
    card_h = 6.42 - top
    big_card(sl, top, card_h)
    steps = d["steps"]
    callout = d.get("callout")
    note = d.get("note")
    n = len(steps)
    x0, x1 = 1.5, 12.1
    gap = 0.42 if n <= 5 else 0.34
    cw = (x1 - x0 - (n - 1) * gap) / n
    ch = 1.5
    cy = top + 0.35
    if not callout and not note:
        cy = top + (card_h - ch) / 2
    elif not callout:
        cy = top + (card_h - ch - 0.9) / 2 + 0.1
    if n <= 3:
        ts = 13
    elif n <= 5:
        ts = 11.5
    elif n <= 6:
        ts = 10.5
    else:
        ts = 9
    ss = max(8, ts - 2.5)
    for i, (st, sub) in enumerate(steps):
        x = x0 + i * (cw + gap)
        rect(sl, x, cy, cw, ch, MINI, radius=0.12)
        rect(sl, x + 0.1, cy - 0.026, cw - 0.2, 0.07, TERRA, radius=0.035)
        paras = [[(st, ts, True, INK, 0)]]
        if sub:
            paras.append([(sub, ss, False, GRAY, 0)])
        text(sl, x + 0.08, cy, cw - 0.16, ch, paras, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05, space_after=3)
        if i < n - 1:
            text(sl, x + cw - 0.06, cy + ch / 2 - 0.22, gap + 0.14, 0.4,
                 [[("→", 18, True, TERRA, 0)]], align=PP_ALIGN.CENTER, wrap=False)
    yy = cy + ch + 0.3
    if callout:
        cw2 = 6.4
        cx2 = (PW - cw2) / 2
        ch2 = 1.15
        rect(sl, cx2, yy, cw2, ch2, MINI, radius=0.12)
        rect(sl, cx2 + 0.12, yy - 0.026, cw2 - 0.24, 0.07, TERRA, radius=0.035)
        text(sl, cx2 + 0.3, yy + 0.12, cw2 - 0.6, ch2 - 0.24,
             [[(callout[0], 12.5, True, INK, 0)], [(callout[1], 10, False, GRAY, 0)]],
             line_spacing=1.1, space_after=3)
    if note:
        two = est_lines((note[0] or "") + " " + note[1], 10.0, 10.5) > 1
        nh = 0.72 if two else 0.5
        note_pill(sl, 1.45, top + card_h - nh - 0.2, 10.65, note[0], note[1], two_line=two)


def t_hub(sl, d):
    bg(sl, CREAM)
    kicker_rule(sl, d["k"])
    top = heading(sl, d["h"])
    footer(sl)
    card_h = 6.42 - top - 0.62
    big_card(sl, top, card_h)
    cx, cy = PW / 2, top + card_h / 2 + 0.08
    # connectors
    for (tx, ty) in ((cx, top + 0.6), (2.6, cy - 0.4), (PW - 2.6, cy - 0.4),
                     (4.6, top + card_h - 0.6), (PW - 4.6, top + card_h - 0.6)):
        ln = slide_line(sl, cx, cy, tx, ty)
    # center circle
    dmt = 2.35
    rect(sl, cx - dmt / 2, cy - dmt / 2, dmt, dmt, MINI, line=TERRA, line_w=1.5, shape=MSO_SHAPE.OVAL)
    ctr = d["center"]
    text(sl, cx - dmt / 2 + 0.15, cy - 0.5, dmt - 0.3, 1.0,
         [[(ctr[0], 13.5, True, INK, 0)], [(ctr[1], 9, False, GRAY, 0)]],
         align=PP_ALIGN.CENTER, line_spacing=1.05, space_after=3)
    # spokes: top, left, right, bottom-left, bottom-right
    sw, sh = 3.0, 1.02
    pos = [(cx - sw / 2, top + 0.22),
           (1.62, cy - sh / 2 - 0.15),
           (PW - 1.62 - sw, cy - sh / 2 - 0.15),
           (3.15, top + card_h - sh - 0.16),
           (PW - 3.15 - sw, top + card_h - sh - 0.16)]
    for (x, y), (st, sub) in zip(pos, d["spokes"]):
        rect(sl, x, y, sw, sh, MINI, radius=0.12)
        rect(sl, x + 0.1, y - 0.024, sw - 0.2, 0.065, TERRA, radius=0.032)
        text(sl, x + 0.2, y + 0.12, sw - 0.4, sh - 0.2,
             [[(st, 11.5, True, INK, 0)], [(sub, 9, False, GRAY, 0)]],
             line_spacing=1.05, space_after=2)
    note = d.get("note")
    if note:
        note_pill(sl, 1.15, top + card_h + 0.12, 11.25, note[0], note[1])


def slide_line(sl, x1, y1, x2, y2):
    from pptx.util import Emu as E
    ln = sl.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = CONNECT
    ln.line.width = Pt(1.0)
    ln.shadow.inherit = False
    return ln


def t_script(sl, d):
    bg(sl, CREAM)
    kicker_rule(sl, d["k"])
    top = heading(sl, d["h"], d.get("sub"))
    footer(sl)
    x0 = 1.15
    colw = 5.5
    x_r = x0 + colw + 0.25
    hh = 0.5
    for x, label in ((x0, d["lh"]), (x_r, d["rh"])):
        rect(sl, x, top, colw, hh, TERRA_SOFT2, radius=0.14)
        text(sl, x + 0.28, top, colw - 0.56, hh, [[(label, 10.5, True, INK, 200)]],
             anchor=MSO_ANCHOR.MIDDLE)
    note = d.get("note")
    body_y = top + hh + 0.14
    bot = 6.42 - (0.6 if note else 0)
    rows = d["rows"]
    gap = 0.12
    rh = (bot - body_y - (len(rows) - 1) * gap) / len(rows)
    for i, (say, tag, how, cue) in enumerate(rows):
        y = body_y + i * (rh + gap)
        rect(sl, x0, y, colw, rh, MINI, radius=0.1)
        num_circle(sl, x0 + 0.2, y + rh / 2 - 0.15, i + 1, d=0.3)
        paras = [[(say, 10.5, True, INK, 0)]]
        if tag:
            paras.append([(tag, 8.5, False, GRAY, 0)])
        text(sl, x0 + 0.65, y + 0.08, colw - 0.9, rh - 0.16, paras, line_spacing=1.03,
             space_after=2, anchor=MSO_ANCHOR.MIDDLE)
        rect(sl, x_r, y, colw, rh, MINI, radius=0.1)
        text(sl, x_r + 0.3, y + 0.08, colw - 0.6, rh - 0.16,
             [[(how, 10.5, True, INK, 0)], [(cue, 9.5, False, GRAY, 0)]],
             line_spacing=1.05, space_after=2, anchor=MSO_ANCHOR.MIDDLE)
    if note:
        note_pill(sl, x0, 6.42 - 0.5, 11.25, note[0], note[1], size=10)


DISPATCH = {
    "title": lambda sl, d: t_title(sl, d),
    "final": lambda sl, d: t_title(sl, d, final=True),
    "quote": t_quote,
    "bullets": t_bullets,
    "keyidea": t_keyidea,
    "grid": t_grid,
    "compare": t_compare,
    "table": t_table,
    "timeline": t_timeline,
    "hub": t_hub,
    "script": t_script,
}

for d in S:
    sl = prs.slides.add_slide(BLANK)
    DISPATCH[d["t"]](sl, d)

out = "Module 13 - Editable.pptx"
prs.save(out)
print(f"saved {out} with {len(prs.slides.slides._sldIdLst)} slides" if False else f"saved {out} with {len(S)} slides")
